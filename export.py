"""Выгрузка заявок в Excel."""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models import local

COLUMNS = (
    ("№", 7),
    ("Дата", 17),
    ("Клиент", 22),
    ("Контакт", 22),
    ("Источник", 14),
    ("UTM source", 14),
    ("UTM campaign", 18),
    ("Тема", 28),
    ("Срочность", 11),
    ("Стадия", 13),
    ("Ответственный", 18),
    ("Сумма, ₽", 13),
    ("Следующий шаг", 17),
    ("Причина отказа", 22),
    ("Согласие на ПД", 17),
    ("Текст заявки", 60),
)


def _local(moment: datetime | None) -> datetime | None:
    return local(moment).replace(tzinfo=None) if moment else None


def leads_xlsx(leads, stage_ru: dict, urgency_ru: dict) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Заявки"
    sheet.append([title for title, _ in COLUMNS])
    for lead in leads:
        assignee = lead.__dict__.get("assignee")
        sheet.append(
            [
                lead.id,
                _local(lead.created_at),
                lead.client_name,
                lead.client_contact,
                lead.source,
                lead.utm_source,
                lead.utm_campaign,
                lead.topic,
                urgency_ru.get(lead.urgency, lead.urgency),
                stage_ru.get(lead.stage, lead.stage),
                assignee.name if assignee else None,
                float(lead.amount) if lead.amount is not None else None,
                _local(lead.due_at),
                lead.lost_reason,
                _local(lead.consent_at),
                lead.text,
            ]
        )

    head = PatternFill("solid", fgColor="E3F0EE")
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True)
        cell.fill = head
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "DD.MM.YYYY HH:MM"
        row[11].number_format = "# ##0"
        row[-1].alignment = Alignment(wrap_text=True, vertical="top")

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
