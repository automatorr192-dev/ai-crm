import asyncio
import io
import json
import os
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from urllib.parse import quote

import httpx
import pytest
from alembic import command
from alembic.config import Config
from openpyxl import load_workbook
from sqlalchemy import update
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

import app as app_module
import backup
import db
import forms
import notify
import outbound
import tgbot
from hub import Hub
from models import Lead, LeadEvent


async def test_manager_adds_a_phone_call_by_hand(database, manager):
    response = manager.post(
        "/leads",
        data={
            "name": "Полина",
            "contact": "+7 999 120-45-67",
            "text": "Звонила про бота",
            "source": "звонок",
        },
        follow_redirects=False,
    )
    assert response.status_code == 303

    [lead] = await db.get_all_leads()
    full = await db.get_lead(lead.id, full=True)
    assert full.source == "звонок"
    assert full.assignee.name == "Ирина"
    assert full.events[0].kind == "created"
    assert full.events[0].author.name == "Ирина"


async def test_viewer_cannot_add_leads(database, watcher):
    assert watcher.get("/leads/new").status_code == 403
    assert watcher.post("/leads", data={"text": "x"}).status_code == 403


async def test_export_is_a_real_excel_file_with_filters(database, boss):
    first, _ = await db.add_lead("Полина", "@polina", "нужен бот", source="сайт")
    await db.add_lead("Олег", "@oleg", "нужна CRM", source="сайт")
    await db.set_stage(first.id, "won")

    response = boss.get("/export/leads.xlsx?stage=won")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]

    sheet = load_workbook(io.BytesIO(response.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "№"
    assert [row[2] for row in rows[1:]] == ["Полина"]
    assert rows[1][9] == "Сделка"


@pytest.fixture
async def team(database):
    boss = await db.create_user("boss", "Владелец", "very-secret", role="admin")
    irina = await db.create_user("irina", "Ирина", "very-secret")
    oleg = await db.create_user("oleg", "Олег", "very-secret")
    mine, _ = await db.add_lead("Клиент Ирины", "@a", "первая")
    theirs, _ = await db.add_lead("Клиент Олега", "@b", "вторая")
    free, _ = await db.add_lead("Ничей", "@c", "третья")
    await db.set_assignee(mine.id, irina.id)
    await db.set_assignee(theirs.id, oleg.id)
    await db.set_access(irina.id, own_only=True, active=True)
    session = await _login("irina")
    return {
        "boss": boss,
        "irina": irina,
        "mine": mine,
        "theirs": theirs,
        "free": free,
        "client": session,
    }


async def _login(login: str):
    from fastapi.testclient import TestClient

    session = TestClient(app_module.app)
    response = session.post(
        "/login", data={"login": login, "password": "very-secret"}, follow_redirects=False
    )
    assert response.status_code == 303
    return session


async def test_own_only_manager_sees_own_and_unassigned(database, team):
    client = team["client"]
    listing = client.get("/leads").text
    assert "Клиент Ирины" in listing
    assert "Ничей" in listing
    assert "Клиент Олега" not in listing
    assert "Клиент Олега" not in client.get("/").text
    assert client.get(f"/leads/{team['theirs'].id}").status_code == 404
    assert (
        client.post(f"/leads/{team['theirs'].id}/stage", data={"stage": "won"}).status_code == 404
    )
    assert "Клиент Олега" not in client.get("/contacts").text

    sheet = load_workbook(io.BytesIO(client.get("/export/leads.xlsx").content)).active
    names = {row[2] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert names == {"Клиент Ирины", "Ничей"}


async def test_live_feed_skips_other_peoples_leads(database, team):
    irina = await db.get_user(team["irina"].id)
    assert Hub._allowed(irina, {"assignee_id": team["irina"].id})
    assert Hub._allowed(irina, {"assignee_id": None})
    assert not Hub._allowed(irina, {"assignee_id": team["boss"].id})
    assert Hub._allowed(irina, {"id": 1})


async def test_owner_switches_access_and_cannot_lock_himself_out(database, boss):
    person = await db.create_user("oleg", "Олег", "very-secret")
    boss.post(f"/team/{person.id}/access", data={"own_only": "1", "active": "1"})
    assert (await db.get_user(person.id)).own_only

    boss.post(f"/team/{person.id}/access", data={})
    assert not (await db.get_user(person.id)).active

    me = await db.get_user_by_login("boss")
    response = boss.post(f"/team/{me.id}/access", data={}, follow_redirects=False)
    assert "error=self" in response.headers["location"]
    assert (await db.get_user(me.id)).active


async def test_only_owner_changes_access(database, manager):
    other = await db.create_user("oleg", "Олег", "very-secret")
    assert manager.post(f"/team/{other.id}/access", data={}).status_code == 403


async def test_closed_access_ends_the_session(database, boss):
    person = await db.create_user("oleg", "Олег", "very-secret")
    session = await _login("oleg")
    await db.set_access(person.id, own_only=False, active=False)
    assert session.get("/", follow_redirects=False).status_code == 303


async def test_public_form_needs_consent(database, client):
    response = client.post("/api/public/lead", json={"text": "нужен бот"})
    assert response.status_code == 422
    assert await db.get_all_leads() == []


async def test_consent_and_utm_are_saved_and_reported(database, client, boss):
    client.post(
        "/api/public/lead",
        json={
            "text": "нужен бот",
            "consent": True,
            "utm_source": "yandex",
            "utm_medium": "cpc",
            "utm_campaign": "brand",
        },
    )
    [lead] = await db.get_all_leads()
    assert lead.consent_at is not None
    assert (lead.utm_source, lead.utm_medium, lead.utm_campaign) == ("yandex", "cpc", "brand")

    await db.set_amount(lead.id, 50000)
    await db.set_stage(lead.id, "won")
    report = await db.report(30)
    assert report["by_campaign"] == [("yandex", "brand", 1, 1, 50000)]
    assert "brand" in boss.get("/report").text


@pytest.fixture
def hook(monkeypatch):
    monkeypatch.setattr(app_module, "INTAKE_TOKEN", "t0ken")
    return "/hook/t0ken"


async def test_tilda_test_request_gets_ok_and_saves_nothing(database, client, hook):
    response = client.post(hook, data={"test": "test"})
    assert response.status_code == 200
    assert response.text == "ok"
    assert await db.get_all_leads() == []


async def test_tilda_form_becomes_a_lead_with_utm_from_cookies(database, client, hook):
    cookie = "TILDAUTM=" + quote("utm_source=vk|||utm_medium=social|||utm_campaign=autumn")
    response = client.post(
        hook,
        data={
            "Name": "Анна",
            "Phone": "+7 (999) 111-22-33",
            "Comments": "Хочу CRM для салона",
            "Checkbox": "yes",
            "tranid": "123:456",
            "formid": "form789",
            "COOKIES": cookie,
        },
    )
    assert response.text == "ok"
    [lead] = await db.get_all_leads()
    assert (lead.client_name, lead.client_contact, lead.source) == (
        "Анна",
        "+7 (999) 111-22-33",
        "tilda",
    )
    assert lead.text == "Хочу CRM для салона"
    assert (lead.utm_source, lead.utm_campaign) == ("vk", "autumn")
    assert lead.consent_at is not None


async def test_contact_form_7_json_is_understood(database, client, hook):
    response = client.post(
        hook,
        json={
            "your-name": "Олег",
            "your-email": "oleg@mail.ru",
            "your-subject": "Вопрос",
            "your-message": "Сколько стоит?",
        },
    )
    assert response.json()["status"] == "ok"
    [lead] = await db.get_all_leads()
    assert (lead.client_name, lead.client_contact, lead.source) == (
        "Олег",
        "oleg@mail.ru",
        "wordpress",
    )


async def test_unknown_fields_are_kept_in_the_text():
    parsed = forms.parse(
        {"name": "Ира", "phone": "89991112233", "email": "i@ya.ru", "Город": "Казань"}
    )
    assert "город: Казань" in parsed["text"]
    assert "почта: i@ya.ru" in parsed["text"]
    assert parsed["contact"] == "89991112233"


async def test_hook_with_wrong_token_does_not_exist(database, client, hook):
    assert client.post("/hook/guess", data={"Name": "x"}).status_code == 404


async def test_hook_is_off_without_token(database, client, monkeypatch):
    monkeypatch.setattr(app_module, "INTAKE_TOKEN", "")
    assert client.post("/hook/", data={"Name": "x"}).status_code == 404
    assert client.post("/hook/anything", data={"Name": "x"}).status_code == 404


@pytest.fixture
def telegram(monkeypatch):
    sent = []

    async def fake_send(chat_id, text, html=False, button=None):
        sent.append({"chat": chat_id, "text": text, "button": button})
        return True

    monkeypatch.setattr(notify, "CHATS", [-100])
    monkeypatch.setattr(notify, "PUBLIC_URL", "https://crm.example.ru")
    monkeypatch.setattr(tgbot, "enabled", lambda: True)
    monkeypatch.setattr(tgbot, "send", fake_send)
    return sent


async def test_new_lead_reaches_the_team_chat(database, telegram):
    lead, _ = await db.add_lead("<b>Хакер</b>", "@x", "текст <script>", source="сайт")
    await db.set_markup(lead.id, "бот", "high", "черновик")
    await app_module.deliver(lead.id)

    [message] = telegram
    assert message["chat"] == -100
    assert "&lt;b&gt;Хакер&lt;/b&gt;" in message["text"]
    assert "&lt;script&gt;" in message["text"]
    assert "🔴" in message["text"]
    assert message["button"] == ("Открыть в CRM", f"https://crm.example.ru/leads/{lead.id}")


async def test_manual_leads_do_not_ping_the_chat(database, telegram):
    lead, _ = await db.add_lead(None, None, "звонок")
    await app_module.deliver(lead.id, announce=False)
    assert telegram == []


async def test_due_reminder_fires_once_and_rearms_on_new_date(database, telegram):
    lead, _ = await db.add_lead("Полина", "@p", "перезвонить")
    await db.set_due(lead.id, datetime.now(UTC) - timedelta(minutes=1))
    future, _ = await db.add_lead("Олег", "@o", "потом")
    await db.set_due(future.id, datetime.now(UTC) + timedelta(hours=2))

    assert await notify.remind_due() == 1
    assert "Срок по сделке" in telegram[0]["text"]
    assert await notify.remind_due() == 0

    await db.set_due(lead.id, datetime.now(UTC) - timedelta(seconds=5))
    assert await notify.remind_due() == 1


async def test_closed_deals_are_not_reminded(database, telegram):
    lead, _ = await db.add_lead("Полина", "@p", "перезвонить")
    await db.set_due(lead.id, datetime.now(UTC) - timedelta(minutes=1))
    await db.set_stage(lead.id, "won")
    assert await notify.remind_due() == 0


@pytest.fixture
def crm(monkeypatch):
    calls = []
    replies = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return replies.pop(0)

    monkeypatch.setattr(outbound, "TRANSPORT", httpx.MockTransport(handler))
    monkeypatch.setattr(outbound, "DELAYS", (0, 0))
    monkeypatch.setattr(outbound, "BITRIX_WEBHOOK", "")
    monkeypatch.setattr(outbound, "AMO_SUBDOMAIN", "")
    monkeypatch.setattr(outbound, "AMO_TOKEN", "")
    return {"calls": calls, "replies": replies}


async def test_lead_goes_to_bitrix24_with_utm(database, crm, monkeypatch):
    monkeypatch.setattr(outbound, "BITRIX_WEBHOOK", "https://studio.bitrix24.ru/rest/1/abc")
    crm["replies"].append(httpx.Response(200, json={"result": 3465}))
    lead, _ = await db.add_lead(
        "Полина", "8 (999) 120-45-67", "нужен бот", source="сайт", utm={"source": "yandex"}
    )
    await outbound.push(await db.get_lead(lead.id))

    [call] = crm["calls"]
    assert str(call.url) == "https://studio.bitrix24.ru/rest/1/abc/crm.lead.add.json"
    fields = json.loads(call.content)["fields"]
    assert fields["PHONE"] == [{"VALUE": "+79991204567", "VALUE_TYPE": "WORK"}]
    assert fields["UTM_SOURCE"] == "yandex"
    saved = await db.get_lead(lead.id, full=True)
    assert saved.external_id == "3465"
    assert saved.events[-1].kind == "synced"


async def test_amocrm_gets_deal_contact_and_note(database, crm, monkeypatch):
    monkeypatch.setattr(outbound, "AMO_SUBDOMAIN", "studio")
    monkeypatch.setattr(outbound, "AMO_TOKEN", "long-lived")
    crm["replies"] += [
        httpx.Response(503),
        httpx.Response(200, json=[{"id": 777, "contact_id": 9, "request_id": ["0"]}]),
        httpx.Response(200, json={"_embedded": {"notes": [{"id": 1}]}}),
    ]
    lead, _ = await db.add_lead("Олег", "oleg@mail.ru", "нужна CRM")
    await outbound.push(await db.get_lead(lead.id))

    first, retry, note = crm["calls"]
    assert retry.headers["authorization"] == "Bearer long-lived"
    body = json.loads(retry.content)[0]
    assert body["_embedded"]["contacts"][0]["custom_fields_values"][0]["field_code"] == "EMAIL"
    assert str(note.url).endswith("/api/v4/leads/notes")
    assert json.loads(note.content)[0]["entity_id"] == 777
    assert (await db.get_lead(lead.id)).external_id == "777"


async def test_refused_sync_is_visible_in_history(database, crm, monkeypatch):
    monkeypatch.setattr(outbound, "BITRIX_WEBHOOK", "https://studio.bitrix24.ru/rest/1/abc")
    crm["replies"].append(
        httpx.Response(
            401, json={"error": "INVALID_CREDENTIALS", "error_description": "Неверный вебхук"}
        )
    )
    lead, _ = await db.add_lead(None, None, "текст")
    await outbound.push(await db.get_lead(lead.id))

    saved = await db.get_lead(lead.id, full=True)
    assert saved.external_id is None
    assert saved.events[-1].kind == "sync_failed"
    assert "Неверный вебхук" in saved.events[-1].note


async def test_nothing_is_sent_without_settings(database, crm):
    lead, _ = await db.add_lead(None, None, "текст")
    await outbound.push(await db.get_lead(lead.id))
    assert crm["calls"] == []


async def test_journal_is_for_the_owner(database, boss, manager):
    lead, _ = await db.add_lead("Полина", None, "текст")
    irina = await db.get_user_by_login("irina")
    await db.set_stage(lead.id, "won", irina.id)

    page = boss.get("/log")
    assert page.status_code == 200
    assert "Ирина" in page.text and "Сделка" in page.text
    assert manager.get("/log").status_code == 403


async def test_due_typed_in_moscow_is_stored_in_utc(database, manager):
    lead, _ = await db.add_lead(None, None, "текст")
    manager.post(f"/leads/{lead.id}/due", data={"due_at": "2026-10-07T14:00"})
    saved = await db.get_lead(lead.id)
    due = saved.due_at if saved.due_at.tzinfo else saved.due_at.replace(tzinfo=UTC)
    assert due == datetime(2026, 10, 7, 11, 0, tzinfo=UTC)
    assert saved.due_input == "2026-10-07T14:00"


async def _age(lead_id: int, days: int) -> None:
    moment = datetime.now(UTC) - timedelta(days=days)
    async with db.Session() as session:
        await session.execute(
            update(LeadEvent).where(LeadEvent.lead_id == lead_id).values(created_at=moment)
        )
        await session.execute(update(Lead).where(Lead.id == lead_id).values(created_at=moment))
        await session.commit()


async def test_forgotten_deals_are_filtered_out(database, boss):
    forgotten, _ = await db.add_lead("Забытый", "@old", "давно")
    await _age(forgotten.id, 5)
    lively, _ = await db.add_lead("Живой", "@new", "вчера")
    closed, _ = await db.add_lead("Закрытый", "@won", "давно")
    await db.set_stage(closed.id, "won")
    await _age(closed.id, 5)

    stale = await db.get_all_leads(stale=True)
    assert [lead.client_name for lead in stale] == ["Забытый"]
    assert "Живой" not in boss.get("/leads?stale=true").text


async def test_morning_digest_lists_overdue_and_forgotten(database, telegram):
    forgotten, _ = await db.add_lead("Забытый", "@old", "давно")
    await _age(forgotten.id, 5)
    late, _ = await db.add_lead("Опоздавший", "@late", "перезвонить")
    await db.set_due(late.id, datetime.now(UTC) - timedelta(hours=2))

    assert await notify.send_digest()
    text = telegram[0]["text"]
    assert "Сводка" in text
    assert "Просрочено: 1" in text and "Опоздавший" in text
    assert "Без движения" in text and "Забытый" in text
    assert telegram[0]["button"][1].endswith("/leads?stale=true")


async def test_backup_restores_into_an_empty_base(database, tmp_path, monkeypatch):
    user = await db.create_user("irina", "Ирина", "very-secret")
    lead, _ = await db.add_lead("Полина", "+7 999 120-45-67", "нужен бот", utm={"source": "vk"})
    await db.set_amount(lead.id, 45000.5, user.id)
    await db.add_note(lead.id, user.id, "созвонились")
    data = await backup.snapshot()

    url = f"sqlite+aiosqlite:///{tmp_path / 'restored.db'}"
    monkeypatch.setenv("DATABASE_URL", url)
    engine = create_async_engine(url)
    monkeypatch.setattr(db, "engine", engine)
    monkeypatch.setattr(db, "Session", async_sessionmaker(engine, expire_on_commit=False))
    config = Config(os.path.join(os.path.dirname(__file__), "..", "alembic.ini"))
    await asyncio.to_thread(command.upgrade, config, "head")

    counts = await backup.restore(data)
    assert counts["leads"] == 1 and counts["notes"] == 1
    restored = await db.get_lead(lead.id, full=True)
    assert restored.amount == Decimal("45000.50")
    assert restored.utm_source == "vk"
    assert restored.notes[0].author.name == "Ирина"
    fresh, _ = await db.add_lead(None, "@next", "новая после восстановления")
    assert fresh.id == lead.id + 1
    await engine.dispose()


async def test_only_owner_downloads_backup(database, boss, manager):
    response = boss.get("/backup")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/gzip"
    assert manager.get("/backup").status_code == 403


async def test_files_are_attached_and_guarded(database, team, monkeypatch, tmp_path):
    monkeypatch.setattr(db, "DATA_DIR", str(tmp_path))
    boss = await _login("boss")
    boss.post(
        f"/leads/{team['theirs'].id}/files",
        files={"file": ("смета.pdf", b"%PDF-1.4 test", "application/pdf")},
    )
    saved = await db.get_lead(team["theirs"].id, full=True)
    [item] = saved.attachments
    assert item.name == "смета.pdf" and item.size == 13
    assert saved.events[-1].kind == "file"

    download = boss.get(f"/files/{item.id}")
    assert download.content == b"%PDF-1.4 test"
    assert download.headers["content-type"] == "application/octet-stream"
    assert "attachment" in download.headers["content-disposition"]
    assert team["client"].get(f"/files/{item.id}").status_code == 404


async def test_oversized_file_is_refused(database, manager, monkeypatch, tmp_path):
    monkeypatch.setattr(db, "DATA_DIR", str(tmp_path))
    monkeypatch.setattr(app_module, "MAX_FILE_BYTES", 10)
    lead, _ = await db.add_lead(None, None, "текст")
    response = manager.post(f"/leads/{lead.id}/files", files={"file": ("big.bin", b"x" * 11)})
    assert response.status_code == 413
    assert (await db.get_lead(lead.id, full=True)).attachments == []


async def test_new_lead_reaches_max_chat(database, monkeypatch):
    import maxbot

    calls = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        return httpx.Response(200, json={"message": {}})

    monkeypatch.setattr(maxbot, "TOKEN", "max-token")
    monkeypatch.setattr(maxbot, "CHATS", [555])
    monkeypatch.setattr(maxbot, "TRANSPORT", httpx.MockTransport(handler))
    monkeypatch.setattr(notify, "CHATS", [])
    monkeypatch.setattr(notify, "PUBLIC_URL", "https://crm.example.ru")
    monkeypatch.setattr(tgbot, "enabled", lambda: False)

    lead, _ = await db.add_lead("Полина", "@p", "нужен бот", source="сайт")
    await app_module.deliver(lead.id)

    [call] = calls
    assert str(call.url).startswith("https://platform-api2.max.ru/messages?chat_id=555")
    assert call.headers["authorization"] == "max-token"
    sent = json.loads(call.content)
    assert sent["format"] == "html" and "Новая заявка" in sent["text"]
    button = sent["attachments"][0]["payload"]["buttons"][0][0]
    assert button == {
        "type": "link",
        "text": "Открыть в CRM",
        "url": f"https://crm.example.ru/leads/{lead.id}",
    }


def test_max_trusts_the_russian_root_certificate():
    import maxbot

    assert maxbot._trust().cert_store_stats()["x509_ca"] > 100
